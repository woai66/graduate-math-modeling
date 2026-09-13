#!/usr/bin/env python3
"""国奖级论文结构和内容配额预检。"""
from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--min-paper-chars", type=int, default=15000)
    parser.add_argument("--min-question-chars", type=int, default=3000)
    parser.add_argument("--min-abstract-chars", type=int, default=1000)
    parser.add_argument("--min-tables", type=int, default=4)
    parser.add_argument("--min-figures", type=int, default=8)
    args = parser.parse_args()
    if not args.docx.is_file():
        print(f"BLOCKED: file not found: {args.docx}")
        return 2
    doc = Document(args.docx)
    paragraph_objects = [p for p in doc.paragraphs if p.text.strip()]
    paragraphs = [p.text.strip() for p in paragraph_objects]
    text = "".join(paragraphs)
    required = ["摘要", "问题重述", "问题一", "问题二", "问题三", "问题四", "模型评价", "参考文献", "附录"]
    errors = [f"missing heading/content: {item}" for item in required if item not in text]
    abstract_match = re.search(r"摘要(.*?)(?:关键词|关键字)", text)
    abstract_chars = len(abstract_match.group(1)) if abstract_match else 0
    if abstract_chars < args.min_abstract_chars:
        errors.append(f"abstract chars={abstract_chars} < {args.min_abstract_chars}")
    question_lengths = {}
    ordinal_map = {"一": 1, "二": 2, "三": 3, "四": 4}
    heading_items = [(index, paragraph.text.strip()) for index, paragraph in enumerate(paragraph_objects) if paragraph.style.name.startswith("Heading") and re.search(r"(?:\d+\s*)?问题[一二三四]", paragraph.text)]
    for index, (paragraph_index, heading_text) in enumerate(heading_items):
        ordinal_match = re.search(r"问题([一二三四])", heading_text)
        if not ordinal_match:
            continue
        ordinal = ordinal_match.group(1)
        number = ordinal_map[ordinal]
        end_index = heading_items[index + 1][0] if index + 1 < len(heading_items) else len(paragraphs)
        section = "".join(paragraphs[paragraph_index + 1:end_index])
        question_lengths[number] = len(section)
        if len(section) < args.min_question_chars:
            errors.append(f"question {number} chars={len(section)} < {args.min_question_chars}")
        required_terms = ["分析", "数据", "模型", "结果", "验证", "敏感", "工程解释", "问题小结"]
        missing_terms = [term for term in required_terms if term not in section]
        if missing_terms:
            errors.append(f"question {number} missing closure terms: {','.join(missing_terms)}")
    with zipfile.ZipFile(args.docx) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    figures = len(root.findall(".//w:drawing", NS))
    print(f"chars={len(text)}; tables={len(doc.tables)}; figures={figures}; question_chars={question_lengths}")
    if len(text) < args.min_paper_chars:
        errors.append(f"paper chars={len(text)} < {args.min_paper_chars}")
    if len(doc.tables) < args.min_tables:
        errors.append(f"tables={len(doc.tables)} < {args.min_tables}")
    if figures < args.min_figures:
        errors.append(f"figures={figures} < {args.min_figures}")
    if errors:
        print("BLOCKED: award-level structure check failed")
        for error in errors:
            print(f"- {error}")
        return 1
    print("PASS: award-level structure quota")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
