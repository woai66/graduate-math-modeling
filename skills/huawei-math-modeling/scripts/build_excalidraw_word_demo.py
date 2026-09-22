"""生成 A4 论文版心内的 Excalidraw 插图检查件；输入为实际插件导出的 PNG。"""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

ASSETS = Path(__file__).resolve().parents[1] / 'assets'
STEM = 'excalidraw-award-style-demo'
image_path = ASSETS / f'{STEM}.png'
source_path = ASSETS / f'{STEM}.excalidraw'
output_path = ASSETS / f'{STEM}-word.docx'


def font(style, size, east_asia='宋体', latin='Times New Roman'):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), east_asia)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    font(doc.styles['Normal'], 12)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.25
    normal.space_after = Pt(6)
    font(doc.styles['Title'], 16, '黑体')
    doc.styles['Title'].paragraph_format.space_before = Pt(0)
    doc.styles['Title'].paragraph_format.space_after = Pt(10)
    doc.styles['Title'].paragraph_format.keep_with_next = True
    font(doc.styles['Caption'], 10.5)
    doc.styles['Caption'].paragraph_format.line_spacing = 1
    doc.styles['Caption'].paragraph_format.space_before = Pt(5)
    doc.styles['Caption'].paragraph_format.space_after = Pt(0)
    title = doc.add_paragraph('园林美学技术路线图排版样例', style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 清除模板和 WPS 继承的标题底线，保持论文式纯文字标题。
    for element in [doc.styles['Title'].element, title._p]:
        ppr = element.get_or_add_pPr()
        for old in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(old)
        border = OxmlElement('w:pBdr')
        for edge in ('top', 'bottom', 'left', 'right', 'between'):
            line = OxmlElement(f'w:{edge}')
            line.set(qn('w:val'), 'nil')
            border.append(line)
        ppr.append(border)
    p = doc.add_paragraph('本页用于检查技术路线图在 A4 论文中的呈现效果。图中以园林美学为例，展示数据预处理、分问题建模与综合评价之间的关系。')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.keep_with_next = True
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.line_spacing = 1
    figure.paragraph_format.space_before = Pt(0)
    figure.paragraph_format.space_after = Pt(0)
    figure.paragraph_format.keep_with_next = True
    pic = figure.add_run().add_picture(str(image_path), width=Cm(16))
    pic._inline.docPr.set('descr', '园林美学技术路线图。上方为数据预处理，中部并列趣味性、幻境感与相似度建模，右侧泛化验证通过反馈箭头连接前两问，底部汇总美学特征。')
    caption = doc.add_paragraph('图 1  园林美学研究技术路线示例', style='Caption')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = '园林美学技术路线图排版样例'
    doc.core_properties.subject = 'Excalidraw 可编辑图形与论文插图测试'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.save(output_path)
    with Image.open(image_path) as im:
        width, height = im.size
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    record = {
        'source': source_path.name, 'source_sha256': sha(source_path),
        'image': image_path.name, 'image_sha256': sha(image_path),
        'docx': output_path.name, 'docx_sha256': sha(output_path),
        'image_pixels': [width, height], 'placement_width_cm': 16,
        'placement_height_cm': 16 * height / width,
        'effective_dpi': width / (16 / 2.54),
        'min_label_point_size': 24 * (16 / 2.54 * 72) / 1128,
        'render_review': 'pending',
    }
    (ASSETS / f'{STEM}.word-build.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    build()
