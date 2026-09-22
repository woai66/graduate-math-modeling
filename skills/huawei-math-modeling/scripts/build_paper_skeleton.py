#!/usr/bin/env python3
"""以官方论文模板为底稿，生成可直接填写的研赛论文 DOCX 骨架。

封面、摘要页的版式（官方 logo、报名表格、华文新魏/隶书标题、页边距、页码起始 0）
全部沿用官方附件3 模板本身；本脚本负责：
  1. 覆盖正文样式为附件2 要求：题目三号黑体、一级标题四号黑体、其余小四宋体、单倍行距；
  2. 在摘要页插入摘要/创新点/关键词的填写占位；
  3. 插入「目录」页与真正的 TOC 域（可在 Word 中按 F9 更新）；
  4. 追加优秀论文范式的正文骨架：一级标题用中文序号，二级标题用 N.M。

为什么以模板为底稿而不是重新绘制封面：官方 logo 尺寸、表格线、标题字体与行距都在模板里，
重建必然产生偏差。

用法：
    python build_paper_skeleton.py <输出.docx> [--template <官方模板.docx>]
                                   [--questions 4] [--title 论文题目] [--no-toc]

生成的 TOC 域初始为空，需要在 Word 中更新一次；可用 --template 之外的
scripts/update_word_fields 步骤完成（本仓库在生成后由 Word 更新并另存）。
"""
from __future__ import annotations

import argparse
import glob
import os
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import (
        WD_ALIGN_PARAGRAPH,
        WD_LINE_SPACING,
        WD_TAB_ALIGNMENT,
        WD_TAB_LEADER,
    )
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    print("BLOCKED: 缺少 python-docx，请先执行 uv add python-docx")
    raise SystemExit(2)


SONG = "宋体"
HEI = "黑体"
WEST = "Times New Roman"
MONO = "Consolas"

PT_HEAD_BIG = 16.0   # 三号：题目、目录标题
PT_HEAD = 14.0       # 四号：一级标题
PT_BODY = 12.0       # 小四：正文、目录条目
PT_CAPTION = 10.5    # 五号：图表题注、参考文献

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x59, 0x59, 0x59)

THEME_ATTRS = ("w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme", "w:cstheme", "w:cs")

CN_NUM = "一二三四五六七八九十"

TEXT_WIDTH_CM = 21.0 - 2.25 - 2.25   # 版心宽度 16.5 cm

DEFAULT_TEMPLATE_GLOBS = [
    r"F:\goodstudy\研0\竞赛\数模\附件3*.docx",
    r"F:\goodstudy\研0\竞赛\数模\**\附件3*.docx",
]


def find_template() -> Path | None:
    env = os.environ.get("MATHMODEL_TEMPLATE")
    if env and Path(env).is_file():
        return Path(env)
    for pattern in DEFAULT_TEMPLATE_GLOBS:
        for hit in sorted(glob.glob(pattern, recursive=True)):
            if Path(hit).is_file():
                return Path(hit)
    return None


def set_rpr_fonts(rpr, east_asian: str, ascii_font: str) -> None:
    """设置字体并清除主题属性，否则 Word 会用主题字体覆盖黑体/宋体。"""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in THEME_ATTRS:
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:eastAsia"), east_asian)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


def style_fonts(style, east_asian, ascii_font, size_pt, color=BLACK, bold=None) -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color
    style.font.bold = bold
    set_rpr_fonts(style.element.get_or_add_rPr(), east_asian, ascii_font)


def run_fonts(run, east_asian, ascii_font, size_pt, color=BLACK, bold=None) -> None:
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    set_rpr_fonts(run._r.get_or_add_rPr(), east_asian, ascii_font)


def set_indent_chars(style, chars: int) -> None:
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if chars:
        ind.set(qn("w:firstLineChars"), str(chars * 100))
        ind.set(qn("w:firstLine"), "240")
    else:
        for attr in ("w:firstLineChars", "w:firstLine"):
            if ind.get(qn(attr)) is not None:
                del ind.attrib[qn(attr)]


def set_left_indent_chars(style, chars: int) -> None:
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:leftChars"), str(chars * 100))
    ind.set(qn("w:left"), str(int(chars * 240)))


def clear_all_indents(style) -> None:
    """清空样式上的全部缩进。

    官方模板里 Caption 等样式带着历史遗留的左缩进（Caption 为 3 cm），
    不清掉会出现“居中但整体偏右”的现象。
    """
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return
    for attr in (
        "w:firstLineChars", "w:firstLine", "w:hangingChars", "w:hanging",
        "w:leftChars", "w:left", "w:rightChars", "w:right", "w:start", "w:end",
    ):
        if ind.get(qn(attr)) is not None:
            del ind.attrib[qn(attr)]


def clear_para_indent(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:firstLine"), "0")


def clear_table_indents(doc) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    clear_para_indent(para)


def trim_trailing_empty_paragraphs(anchor_para) -> int:
    """删除锚点之后、正文之前的空段落（模板尾部遗留），避免产生空白页。"""
    removed = 0
    nxt = anchor_para._element.getnext()
    while nxt is not None and nxt.tag == qn("w:p"):
        following = nxt.getnext()
        para = Paragraph(nxt, anchor_para._parent)
        has_image = bool(nxt.findall(".//" + qn("a:blip")))
        if not para.text.strip() and not has_image:
            nxt.getparent().remove(nxt)
            removed += 1
        else:
            break
        nxt = following
    return removed


def clear_spacing(style) -> None:
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_right_dot_tab(style) -> None:
    """右对齐 + 点线引导，目录页码才能像样例那样排到行尾。"""
    tabs = style.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def set_cell_bottom_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    tc_pr.append(borders)


def add_toc_field(paragraph) -> None:
    """插入真正的 TOC 域，Word 中可按 F9 更新。"""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中按 Ctrl+A 后按 F9，或右键选择“更新域”生成目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, placeholder, end):
        run._r.append(el)
    run_fonts(run, SONG, WEST, PT_BODY)


def override_styles(doc) -> None:
    normal = doc.styles["Normal"]
    style_fonts(normal, SONG, WEST, PT_BODY)
    clear_spacing(normal)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_indent_chars(normal, 2)

    for name, size in (("Heading 1", PT_HEAD), ("Heading 2", PT_BODY), ("Heading 3", PT_BODY)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        style_fonts(st, HEI, WEST, size)
        clear_spacing(st)
        clear_all_indents(st)
        st.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
        )
        st.paragraph_format.keep_with_next = True

    def ensure(name, ea, west, size, align, indent=0, color=BLACK, bold=None, left_chars=0):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, 1)
        st.base_style = doc.styles["Normal"]
        clear_all_indents(st)
        clear_spacing(st)
        style_fonts(st, ea, west, size, color, bold)
        st.paragraph_format.alignment = align
        if indent:
            set_indent_chars(st, indent)
        if left_chars:
            set_left_indent_chars(st, left_chars)
        return st

    ensure("Caption", SONG, WEST, PT_CAPTION, WD_ALIGN_PARAGRAPH.CENTER)
    ensure("目录标题", HEI, WEST, PT_HEAD_BIG, WD_ALIGN_PARAGRAPH.CENTER)
    ensure("表格文字", SONG, WEST, PT_BODY, WD_ALIGN_PARAGRAPH.CENTER)
    ref = ensure("参考文献条目", SONG, WEST, PT_CAPTION, WD_ALIGN_PARAGRAPH.LEFT, left_chars=2)
    ref.paragraph_format.first_line_indent = Pt(-21)
    ensure("代码", SONG, MONO, PT_CAPTION, WD_ALIGN_PARAGRAPH.LEFT)
    ensure("提示", SONG, WEST, PT_CAPTION, WD_ALIGN_PARAGRAPH.LEFT, color=GRAY)

    # 目录条目样式：一级加粗，二级/三级逐级缩进，均带点线引导。
    # Word 内置样式名为小写 "toc 1"…，此处把大小写两种都覆盖，避免设置落空。
    for level in (1, 2, 3):
        for name in (f"TOC {level}", f"toc {level}"):
            try:
                st = doc.styles[name]
            except KeyError:
                if not name.startswith("TOC"):
                    continue
                st = doc.styles.add_style(name, 1)
            st.base_style = doc.styles["Normal"]
            clear_all_indents(st)
            clear_spacing(st)
            style_fonts(st, SONG, WEST, PT_BODY, BLACK, bold=(level == 1))
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if level > 1:
                set_left_indent_chars(st, 2 * (level - 1))
            add_right_dot_tab(st)


def find_paragraph(doc, pattern: str):
    rx = re.compile(pattern)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and rx.search(text):
            return para
    return None


def append_paragraph(doc, text: str, style: str | None = None):
    return doc.add_paragraph(text, style=style)


def heading(doc, text: str, level: int):
    return doc.add_heading(text, level=level)


def guidance(doc, text: str):
    return append_paragraph(doc, text, style="提示")


def caption(doc, text: str):
    return append_paragraph(doc, text, style="Caption")


def cn_num(index: int) -> str:
    if 1 <= index <= len(CN_NUM):
        return CN_NUM[index - 1]
    return str(index)


def build_toc_page(doc) -> None:
    para = append_paragraph(doc, "目  录", style="目录标题")
    para.paragraph_format.page_break_before = True
    toc_para = append_paragraph(doc, "")
    toc_para.paragraph_format.first_line_indent = Pt(0)
    add_toc_field(toc_para)


def build_question(doc, index: int, chapter: int) -> None:
    ordinal = ["一", "二", "三", "四", "五"][index - 1]
    heading(doc, f"{cn_num(chapter)}.问题{ordinal}模型建立与求解", 1)
    guidance(doc, "【本问要回答的任务、题面编号与输出要求】")
    sections = [
        ("问题分析", "【本问与前后问的联系、难点、可计算定义与反例、选模理由】"),
        ("数据与特征", "【使用的数据、字段、单位、缺失与异常处理、特征构造依据】"),
        ("模型建立", "【变量、假设、目标函数或统计模型、公式推导与符号说明】"),
        ("参数估计与求解算法", "【参数来源、初值与搜索范围、停止条件、复杂度、求解器版本】"),
        ("结果与分析", "【真实运行结果、图表编号与解释、与基线的比较】"),
        ("验证与敏感性分析", "【独立验证方式、关键参数扰动、失败样本与适用范围】"),
        ("小结", "【本问结论、单位、不确定性与对下一问的接口】"),
    ]
    for order, (sub_title, hint) in enumerate(sections, start=1):
        heading(doc, f"{chapter}.{order} {sub_title}", 2)
        guidance(doc, hint)
        if sub_title == "结果与分析":
            caption(doc, f"图 {chapter}-1 【图题：说明对象、条件与单位】")
            caption(doc, f"表 {chapter}-1 【表题：说明对象、条件与单位】")


def append_body(doc, questions: int) -> None:
    heading(doc, "一.问题重述", 1)
    heading(doc, "1.1 问题背景", 2)
    guidance(doc, "【只写与本题相关的背景，避免大段科普。】")
    heading(doc, "1.2 问题提出", 2)
    guidance(doc, "【逐问复述任务动词、对象、输入、输出、约束和评价口径。】")

    heading(doc, "二.问题分析与技术路线", 1)
    heading(doc, "2.1 总体思路", 2)
    guidance(doc, "【说明各问之间的依赖关系、跨问接口和不确定性传递。】")
    heading(doc, "2.2 各问难点与建模思路", 2)
    guidance(doc, "【逐问说明难点、候选路线、选择理由和验证方式。】")
    heading(doc, "2.3 技术路线图", 2)
    caption(doc, "图 2-1 【技术路线图：与正文各问一一对应】")

    heading(doc, "三.模型假设与符号说明", 1)
    heading(doc, "3.1 模型假设", 2)
    guidance(doc, "【假设 1：……（写清工程或数据依据，并说明可检验方式）。】")
    heading(doc, "3.2 符号说明", 2)
    table = doc.add_table(rows=3, cols=3)
    data = [("符号", "含义", "单位"), ("x", "决策变量", "—"), ("ρ", "空气密度", "kg/m³")]
    for r, row in enumerate(data):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            for para in cell.paragraphs:
                para.style = doc.styles["表格文字"]
                clear_para_indent(para)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tbl_pr.append(borders)
    for cell in table.rows[0].cells:
        set_cell_bottom_border(cell)
    caption(doc, "表 3-1 符号说明")

    heading(doc, "四.数据来源与预处理", 1)
    heading(doc, "4.1 数据来源与字段说明", 2)
    guidance(doc, "【来源、文件清单与哈希、字段、单位、坐标系、时间范围与业务键。】")
    heading(doc, "4.2 数据初检", 2)
    guidance(doc, "【样本量、缺失、异常、重复、标签分布、覆盖范围。】")
    heading(doc, "4.3 数据预处理", 2)
    guidance(doc, "【清洗规则、处理前后行数、防泄漏的划分方式与其依据。】")

    for i in range(1, questions + 1):
        build_question(doc, i, 4 + i)

    review_chapter = 4 + questions + 1
    heading(doc, f"{cn_num(review_chapter)}.模型评价与推广", 1)
    heading(doc, f"{review_chapter}.1 模型优点", 2)
    guidance(doc, "【逐问说明优点，不要只写统一的一段。】")
    heading(doc, f"{review_chapter}.2 模型不足", 2)
    guidance(doc, "【写清失效条件、未覆盖样本与证据边界。】")
    heading(doc, f"{review_chapter}.3 改进方向与推广", 2)
    guidance(doc, "【说明进一步可做的实验与适用范围。】")

    heading(doc, "参考文献", 1)
    for sample in [
        "[1] 作者，书名，出版地：出版社，起止页码，出版年。",
        "[2] 作者，论文名，杂志名，卷期号：起止页码，出版年。",
        "[3] 作者，资源标题，网址，访问时间（年月日）。",
    ]:
        append_paragraph(doc, sample, style="参考文献条目")
    guidance(doc, "【正文引用处用方括号标注编号，如 [1][3]；引用书籍必须指出页码；每条文献都要被正文引用。】")

    heading(doc, "附录", 1)
    heading(doc, "A.1 人工智能工具使用说明", 2)
    guidance(
        doc,
        "【按附件4 要求列明：工具名称、版本/型号、开发机构/公司、版本发布日期，"
        "以及本队在哪些环节使用、做了哪些核对与修改。】",
    )
    heading(doc, "A.2 核心代码展示", 2)
    guidance(
        doc,
        "【粘贴本问的核心代码，使用“代码”样式；完整源代码按题目要求另行提交竞赛系统。】",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("--template", type=Path, default=None, help="官方论文模板 .docx")
    parser.add_argument("--questions", type=int, default=4, help="题面的问题数量，默认 4")
    parser.add_argument("--title", default=None, help="论文题目；不填则保留模板下划线待填写")
    parser.add_argument("--no-toc", action="store_true", help="不插入目录页")
    args = parser.parse_args()

    if not 1 <= args.questions <= 5:
        print("BLOCKED: --questions 需在 1—5 之间")
        return 2

    template = args.template or find_template()
    if template is None or not Path(template).is_file():
        print("BLOCKED: 找不到官方论文模板。请用 --template 指定附件3 模板 .docx，")
        print("         或设置环境变量 MATHMODEL_TEMPLATE 指向该文件。")
        return 2
    print(f"底稿模板: {template}")

    doc = Document(str(template))
    override_styles(doc)
    clear_table_indents(doc)

    title_para = find_paragraph(doc, r"^题\s*目")
    if title_para is not None:
        for run in list(title_para.runs):
            run._r.getparent().remove(run._r)
        run = title_para.add_run("题    目：")
        run_fonts(run, "隶书", "隶书", 18.0)
        if args.title:
            run = title_para.add_run(args.title)
            run_fonts(run, HEI, WEST, PT_HEAD_BIG)

    kw_para = find_paragraph(doc, r"^关键词")
    abstract_hint = (
        "【摘要需写清：针对问题一……；针对问题二……；最后是总体结论与创新点。"
        "每个定量结论都要有真实运行结果支撑，一般不超过两页，无需英文。】"
    )
    innovation_hint = "【创新点：说明新在哪里、为什么旧方法不够、如何验证、适用边界。】"
    if kw_para is not None:
        kw_para.insert_paragraph_before(abstract_hint, style="提示")
        kw_para.insert_paragraph_before(innovation_hint, style="提示")
        trimmed = trim_trailing_empty_paragraphs(kw_para)
        if trimmed:
            print(f"已清理模板尾部空段落 {trimmed} 个")
    else:
        print("警告: 未找到“关键词”段落，摘要提示未插入")

    if not args.no_toc:
        build_toc_page(doc)

    append_body(doc, args.questions)
    clear_table_indents(doc)

    for probe in (r"^一\.问题重述", r"^1\s*问题重述"):
        first_body = find_paragraph(doc, probe)
        if first_body is not None:
            first_body.paragraph_format.page_break_before = True
            break

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    toc_note = "含目录页（TOC 域，需在 Word 中更新）" if not args.no_toc else "无目录页"
    print(
        f"PASS: 已生成 {args.output}（底稿=官方模板；{args.questions} 个问题章节；{toc_note}）"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
