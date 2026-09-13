from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
RESULTS = PROJECT / "results"
FIGURES = PROJECT / "figures" / "result"
OUTPUT = PROJECT / "paper" / "论文草稿.docx"


def read_csv(name: str):
    with (RESULTS / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def cell(cell, value, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell(t.rows[0].cells[i], h, True, "FFFFFF")
        shade(t.rows[0].cells[i], "315A83")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell(cells[i], value)
            if ri % 2:
                shade(cells[i], "EAF1F7")
    doc.add_paragraph()


def shade(cell_obj, fill):
    props = cell_obj._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p


def figure(doc, name, caption, width=6.0):
    path = FIGURES / name
    if not path.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    r.bold = True
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)


def formula_after(paragraph, text):
    p = OxmlElement("w:p")
    para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    run.append(mt)
    omath.append(run)
    para.append(omath)
    p.append(para)
    paragraph._p.addnext(p)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
    for name in ("Normal", "Body Text"):
        st = doc.styles[name]
        st.font.name = "宋体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        st.font.size = Pt(10.5)
        st.paragraph_format.line_spacing = 1.35
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._r.extend([b, instr, e])


def main():
    q1 = read_csv("q1_candidate_summary.csv")
    q2 = read_csv("q2_cluster_sine_results.csv")
    q3 = read_csv("q3_jrc_results.csv")
    q4 = read_csv("q4_connectivity_results.csv")
    candidates = read_csv("q4_uncertainty_candidates.csv")
    q1_cov = sum(float(r["mask_coverage"]) for r in q1) / len(q1)
    q2_r2 = [float(r["R2"]) for r in q2]
    q2_rmse = [float(r["RMSE_mm"]) for r in q2]
    q3_jrc = [float(r["JRC"]) for r in q3]
    q4_score = [float(r["relative_score"]) for r in q4]
    doc = Document()
    setup(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2025 年中国研究生数学建模竞赛 C 题\n围岩裂隙精准识别与三维模型重构")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = sub.add_run("论文草稿 基于真实附件的增强基线 不可直接提交")
    rr.italic = True
    doc.add_heading("摘要", 1)
    body(doc, f"针对钻孔成像展开图中的裂隙识别、正弦参数表征、粗糙度评价和多钻孔连通性问题，本文按照“数据审计—图像处理—几何拟合—空间推断”的链路建立可复现的增强基线。问题一对附件 1 的 10 张图像采用 CLAHE、双边滤波、黑帽变换、Canny 和形态学处理，平均候选掩膜覆盖率为 {q1_cov:.2%}。问题二对附件 2 进行 DBSCAN 候选簇提取和 RANSAC 线性化正弦拟合，共保留 {len(q2)} 个候选簇，平均 R² 为 {sum(q2_r2)/len(q2_r2):.4f}，平均 RMSE 为 {sum(q2_rmse)/len(q2_rmse):.2f} mm。问题三对附件 3 的 11 张图像提取轮廓并比较采样策略，JRC 基线值范围为 {min(q3_jrc):.2f}—{max(q3_jrc):.2f}。问题四对附件 4 的 7 组相邻钻孔构造未标定的相对连通评分，评分范围为 {min(q4_score):.4f}—{max(q4_score):.4f}。结果表明，增强图像处理可以建立稳定的数据链路，但逐裂隙分离、JRC 工程标定和连通概率校准仍是正式参赛方案必须补齐的环节。")
    body(doc, "关键词：围岩裂隙；图像分割；DBSCAN；RANSAC；JRC；钻孔连通性；三维重构")

    doc.add_heading("1 问题重述与总体思路", 1)
    body(doc, "题目以钻孔成像展开图为数据载体，要求从像素级裂隙识别出发，完成规则裂隙参数化、复杂裂隙粗糙度评价以及多钻孔裂隙网络重构。四个问题具有明确的依赖关系：问题一输出的候选裂隙区域是问题二和问题三轮廓处理的基础，问题二和问题三提供的几何与粗糙度特征又是问题四空间连通性推断的输入。")
    body(doc, "本文先完成不依赖深度学习训练标签的增强基线，再把每一问的结果、局限和后续增强方向分开记录。任何未经过标定的分数均以“相对评分”表述，避免把启发式结果写成工程概率。")
    table(doc, ["问题", "输入", "输出", "主任务"], [
        ["Q1", "附件1，10张图像", "同尺寸候选二值图", "像素级分类识别"],
        ["Q2", "附件2，10张图像", "R、P、β、C及拟合指标", "聚类与几何拟合"],
        ["Q3", "附件3，11张图像", "轮廓、Z₂、JRC和采样比较", "几何统计评价"],
        ["Q4", "附件4，6孔40段图像", "相对连通评分、补孔候选", "图论与空间推断"],
    ])

    doc.add_heading("2 数据审计与模型假设", 1)
    body(doc, "数据审计扫描得到 71 张 JPG 图像：附件 1、2、3 分别为 10、10、11 张，附件 4 为 40 张。附件 1—3 的图像尺寸均为 244×1350，附件 4 的图像尺寸为 864×9167。审计发现附件 1 图 1-2 与附件 3 图 3-3 的 SHA-256 完全相同，后续训练和验证必须按原图分组并记录排除策略。")
    body(doc, "假设一：灰度对比度和梯度可以用于构造裂隙候选区域，但候选区域仍包含纹理和泥浆伪影。假设二：完整正弦状裂隙的周期接近钻孔周长 94.25 mm。假设三：相邻钻孔同深度段的图像特征相似性可以作为连通性排序的弱证据，不能替代几何匹配和现场标定。")
    table(doc, ["符号", "含义", "单位"], [["x", "周向展开坐标", "mm"], ["y", "孔深坐标", "mm"], ["R", "正弦振幅", "mm"], ["P", "正弦周期", "mm"], ["β", "相位", "rad"], ["C", "中心线", "mm"], ["Z₂", "轮廓统计量", "—"], ["JRC", "粗糙度系数", "—"]])

    doc.add_heading("3 问题一 裂隙候选分割", 1)
    q1_method = body(doc, "对灰度图像 I(x,y)，首先使用 CLAHE 增强局部对比度，再用双边滤波削弱纹理噪声。黑帽变换提取暗裂隙结构，Otsu 方法得到自适应阈值；同时用 Canny 提取边缘，二者取并集后进行 3×3 闭运算和开运算。最后使用 8 邻域连通域删除面积小于 12 像素的孤立区域。")
    formula_after(q1_method, "M(x,y) = MorphOpen(MorphClose(BlackHat(I) ∪ Canny(I)))")
    body(doc, f"附件 1 的 10 张图像均成功生成同尺寸候选掩膜，平均覆盖率为 {q1_cov:.2%}。覆盖率范围较宽，说明纹理强度和泥浆伪影会显著影响阈值结果。由于附件中没有独立像素标签，本节不报告 IoU、Dice、Precision、Recall 或 Accuracy。")
    figure(doc, "q1_图1-1_candidate.png", "图 3-1 附件 1 图 1-1 候选裂隙掩膜")
    figure(doc, "q1_coverage.png", "图 3-2 十张附件 1 图像的候选掩膜覆盖率")
    body(doc, "若要进入正式参赛候选稿，应补充人工标注协议和独立验证集，并比较轻量 U-Net、DeepLabV3+ 或其他分割模型；标注集必须按原图或钻孔分组，不能把同一张图的裁剪块随机拆到训练集和验证集。")

    doc.add_heading("4 问题二 正弦状裂隙聚类与鲁棒拟合", 1)
    q2_method = body(doc, "对附件 2，先对平滑图像进行 Canny 边缘检测，将边缘点归一化到 [0,1]² 后使用 DBSCAN 聚类。聚类参数为 eps=0.055、min_samples=20，点数少于 80 的簇被剔除。对每个候选簇，将题面模型线性化：")
    formula_after(q2_method, "y = R sin(2πx/P + β) + C = a sin(2πx/P) + b cos(2πx/P) + C")
    body(doc, "固定 P=94.25 mm 后，使用 RANSAC 回归估计 a、b、C，再由 R=sqrt(a²+b²)、β=atan2(b,a) 得到参数。每幅图最多保留 R² 最高的 3 个候选簇。该过程共保留 17 个候选簇，覆盖 9 幅图像；平均 R² 为 0.1537，平均 RMSE 为 73.89 mm。")
    figure(doc, "q2_r2_boxplot.png", "图 4-1 各图像候选簇正弦拟合 R² 分布")
    table(doc, ["图像", "簇数", "R/mm 均值", "R² 均值", "RMSE/mm 均值"], [[image, str(len([r for r in q2 if r["image"] == image])), f"{sum(float(r['R_mm']) for r in q2 if r['image'] == image)/len([r for r in q2 if r['image'] == image]):.2f}", f"{sum(float(r['R2']) for r in q2 if r['image'] == image)/len([r for r in q2 if r['image'] == image]):.4f}", f"{sum(float(r['RMSE_mm']) for r in q2 if r['image'] == image)/len([r for r in q2 if r['image'] == image]):.2f}"] for image in sorted({r["image"] for r in q2})])
    body(doc, "R² 和 RMSE 表明，当前聚类结果仍混入纹理、拼接线和跨裂隙边缘，不能直接作为题目要求的最终表 1。后续应加入裂隙方向连续性、周期候选搜索、轮廓连通和 RANSAC 内点比例筛选，并对图 2-1 至图 2-3 给出逐裂隙可视化。")

    doc.add_heading("5 问题三 复杂裂隙粗糙度评价", 1)
    q3_method = body(doc, "对附件 3 图像先进行高斯平滑和 Canny 轮廓提取，选择长度不小于 30 的轮廓并取面积最大的轮廓作为候选。为避免重复横坐标造成零间距，先对相同 x 坐标的点取纵坐标中位数，再将像素坐标映射到 94.25 mm 周长和 500 mm 深度。")
    formula_after(q3_method, "Z₂ = sqrt((1/N) Σ[((yᵢ₊₁−yᵢ)/(xᵢ₊₁−xᵢ))²])，JRC = 51.85Z₂^0.6 − 10.37")
    body(doc, f"附件 3 的 11 张图像均提取到候选轮廓，共生成 44 行采样结果。当前 JRC 基线值范围为 {min(q3_jrc):.2f}—{max(q3_jrc):.2f}，且不同采样点数和轮廓形态会造成明显变化。该范围是图像候选轮廓的统计结果，尚未经过工程轮廓标定，不能直接对应 Barton 标准等级。")
    figure(doc, "q3_jrc_sensitivity.png", "图 5-1 图 3-1 的 JRC 采样敏感性")
    table(doc, ["采样方法", "样本数", "JRC 最小值", "JRC 最大值"], [[method, str(len([r for r in q3 if r["method"] == method])), f"{min(float(r['JRC']) for r in q3 if r['method'] == method):.2f}", f"{max(float(r['JRC']) for r in q3 if r['method'] == method):.2f}"] for method in sorted({r["method"] for r in q3})])
    body(doc, "正式方案需要对轮廓进行曲率约束和平滑参数标定，比较等间距、曲率加权和自适应采样，并对面积、采样密度和噪声扰动进行联合敏感性分析。")

    doc.add_heading("6 问题四 多钻孔连通性与补孔候选", 1)
    q4_method = body(doc, "附件 4 的六个钻孔按题面坐标布置。对每个分段计算平滑 Canny 边缘密度、平均亮度和竖直梯度特征。对于相邻钻孔的相同深度段，令特征差异为 d，并定义相对连通评分 p=exp(−50d)。该 p 仅用于排序，未经过真实连通标签校准。")
    formula_after(q4_method, "p_rel = exp(−50d),   d = |Δedge| + 0.002|Δintensity|/255")
    body(doc, f"七组相邻孔对均生成了同深度评分，评分范围为 {min(q4_score):.4f}—{max(q4_score):.4f}。按最接近 0.5 的相对不确定性排序，前三个补孔候选为：{', '.join(f'{r["between_holes"]} 段 {r["segment_start_m"]}m' for r in candidates)}。")
    figure(doc, "q4_connectivity_3d.png", "图 6-1 基于相对评分的六孔三维网络示意")
    table(doc, ["孔间位置", "分段起点/m", "相对不确定性"], [[r["between_holes"], r["segment_start_m"], f"{float(r["uncertainty_score"]):.4f}"] for r in candidates])
    body(doc, "该结果没有使用裂隙的三维姿态、相位、JRC 或真实连通样本，因此只能作为 Q4 的数据链路和可视化基线。正式模型应将单孔裂隙映射为三维节点，使用距离、方向差、深度差、JRC 和图像置信度建立可校准概率模型，再以不确定性下降或覆盖收益/成本比选择补孔位置。")

    doc.add_heading("7 模型评价与后续工作", 1)
    body(doc, "本文完成了从原始附件到结果表、图表和论文草稿的可复现链路，所有实验代码和 CSV 结果均保存在项目目录。增强图像处理降低了输入噪声，DBSCAN 与 RANSAC 提供了逐簇拟合框架，重复横坐标修正避免了 JRC 数值计算中的零间距错误，三维图把钻孔空间布局和评分结果可视化。")
    body(doc, "当前限制包括：缺少独立像素标签，无法完成监督分割评价；Q2 的候选簇仍包含大量纹理，R² 和 RMSE 不足以支撑最终表征；Q3 的轮廓和 JRC 依赖图像边缘，尚无工程标定；Q4 的 p_rel 不是概率，补孔位置没有成本和现场可行性验证。上述限制决定了本文只能作为论文草稿和建模记录。")
    body(doc, "下一轮工作应先完成 Q1 标注协议和 Q2 逐裂隙轮廓分离，再冻结全量结果；随后补齐验证、敏感性、消融、证据矩阵和官方提交规则，最后才进入参赛候选稿审查。")

    doc.add_heading("参考文献", 1)
    body(doc, "[1] 康红普，姜鹏飞，王子越，等. 煤巷钻锚一体化快速掘进技术与装备及应用[J]. 煤炭学报，2024，49(1):131-151。")
    body(doc, "[2] 袁亮，张平松. 煤矿透明地质模型动态重构的关键技术与路径思考[J]. 煤炭学报，2023，48(1):1-14。")
    body(doc, "[3] 2025 年中国研究生数学建模竞赛 C 题题面及附件说明，竞赛资料目录提供的原始文件。")

    doc.add_heading("附录 代码与复现", 1)
    body(doc, "环境由 uv 管理，依赖锁定在项目根目录 uv.lock。主要入口为 projects/2025-C/code/q1_segmentation.py、q2_sine_fit.py、q3_jrc.py、q4_connectivity.py 和 make_project_figures.py。输入数据包不写入仓库，运行前需解压到 MASTER_PROMPT.md 指定的临时目录。")
    body(doc, "复现命令：uv run python projects/2025-C/code/q1_segmentation.py；uv run python projects/2025-C/code/q2_sine_fit.py；uv run python projects/2025-C/code/q3_jrc.py；uv run python projects/2025-C/code/q4_connectivity.py；uv run python projects/2025-C/code/make_project_figures.py。")
    body(doc, "本稿不满足正式参赛候选稿门禁。提交前必须补齐官方规则快照、逐问结果契约、证据矩阵、独立技术审查、独立论文审查和提交附件包。")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
