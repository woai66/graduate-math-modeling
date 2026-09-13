from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
DATA = PROJECT / "results" / "paper_data.json"
FIG = PROJECT / "figures" / "paper"
OUT = PROJECT / "paper" / "国奖范式论文草稿.docx"


def load():
    return json.loads(DATA.read_text(encoding="utf-8"))


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def set_cell(cell, value, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True)
        shade(table.rows[0].cells[i], "315A83")
        table.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if ri % 2:
                shade(cells[i], "EAF1F7")
    doc.add_paragraph()


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_formula_after(paragraph, text):
    p = OxmlElement("w:p")
    para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    run.append(mt)
    omath.append(run)
    para.append(omath)
    p.append(para)
    paragraph._p.addnext(p)


def add_fig(doc, name, caption, width=6.1):
    path = FIG / name
    if not path.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.bold = True
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
    for name in ("Normal", "Body Text"):
        st = doc.styles[name]
        st.font.name = "宋体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        st.font.size = Pt(11)
        st.paragraph_format.line_spacing = 1.45
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def main():
    data = load()
    q1, q2, q3, q4 = data["q1"], data["q2"], data["q3"], data["q4"]
    doc = Document(); setup(doc)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2025 年中国研究生数学建模竞赛 C 题\n围岩裂隙精准识别与三维模型重构")
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("国奖范式结构化论文草稿  基于真实运行结果  不可直接提交").italic = True

    doc.add_heading("摘要", 1)
    add_body(doc, f"围岩裂隙的识别、定量表征和空间重构是煤矿巷道安全评价的重要基础。针对题目给出的四类任务，本文建立从数据审计、图像预处理、裂隙候选提取、几何参数拟合到多钻孔空间推断的统一流程。数据审计扫描 71 张图像，确认附件 1—3 的尺寸为 244×1350，附件 4 的尺寸为 864×9167，并发现附件 1 图 1-2 与附件 3 图 3-3 完全重复。问题一采用 CLAHE、双边滤波、黑帽变换、Otsu、Canny 和形态学方法生成 10 张同尺寸候选掩膜，平均覆盖率为 {q1['mean_cov']:.2f}%。问题二使用 DBSCAN 和 RANSAC 对候选边缘进行聚类与正弦拟合，共保留 {q2['n_clusters']} 个候选簇，平均 R² 为 {q2['mean_r2']:.4f}。问题三对 11 张图像提取复杂轮廓并比较采样方案，JRC 基线范围为 {q3['range_all'][0]:.2f}—{q3['range_all'][1]:.2f}。问题四对七组相邻钻孔对构造相对连通评分，范围为 {q4['range'][0]:.4f}—{q4['range'][1]:.4f}。由于附件未提供独立像素标签，问题一的监督精度尚未具备；由于 Q2—Q4 尚缺工程标定和现场验证，本文定位为国奖范式结构化草稿，所有限制均在文中明确列出。")
    add_body(doc, "关键词：围岩裂隙；图像处理；正弦拟合；粗糙度系数；钻孔网络；不确定性分析")

    doc.add_heading("目录", 1)
    add_bullets(doc, ["1 问题重述与总体技术路线", "2 数据观察与预处理", "3 模型假设与符号说明", "4 问题一 裂隙像素识别", "5 问题二 正弦状裂隙定量分析", "6 问题三 复杂裂隙粗糙度评价", "7 问题四 多钻孔连通性与三维重构", "8 模型验证、敏感性与工程解释", "9 模型评价、创新与推广", "参考文献", "附录 代码与复现说明"])

    doc.add_heading("1 问题重述与总体技术路线", 1)
    add_body(doc, "题目围绕钻孔成像展开图中的裂隙网络提出四个递进问题。第一问要求把裂隙像素从岩石纹理、泥浆和钻进痕迹中分离出来；第二问要求对规则的正弦状裂隙进行聚类和参数表征；第三问要求对复杂轮廓计算粗糙度并分析离散采样影响；第四问要求把六个钻孔的信息统一到空间坐标中，估计相邻裂隙的连通性并选择补充钻孔位置。")
    add_body(doc, "总体技术路线不是四个互相独立的算法，而是一条逐级传递的不确定性链路。Q1 的候选区域影响 Q2、Q3 的轮廓；Q2 的相位、中心线和周期为 Q4 的几何匹配提供先验；Q3 的 JRC 为连通性风险提供补充特征；Q4 将这些结果汇总为可解释的空间网络。")
    add_fig(doc, "fig_route.png", "图 1-1 四问统一技术路线图")
    add_body(doc, "在每个阶段，本文都保留基线、结果和失败边界。基线用于验证输入输出接口，增强模型只有在对比结果和稳定性证据支持时才进入候选方案。")

    doc.add_heading("2 数据观察与预处理", 1)
    add_body(doc, "数据审计首先记录每个文件的相对路径、SHA-256、图像尺寸、灰度分布和 Canny 边缘密度。附件 1、2、3 是单孔 500 mm 深度的展开图，附件 4 是六个钻孔的分段图像。重复图像的存在说明跨附件随机切分会造成信息泄漏，因此后续验证必须按原图和钻孔分组。")
    add_table(doc, ["附件", "样本数", "尺寸", "主要用途"], [["附件1", "10", "244×1350", "像素候选分割"], ["附件2", "10", "244×1350", "正弦裂隙拟合"], ["附件3", "11", "244×1350", "复杂轮廓与 JRC"], ["附件4", "40", "864×9167", "跨孔连通与重构"]])
    add_body(doc, f"附件统计结果显示，附件 1 的候选掩膜覆盖率范围为 {q1['cov_min']:.2f}%—{q1['cov_max']:.2f}%，附件 2 的正弦候选簇拟合 R² 存在明显的图像间差异，附件 3 的 JRC 对采样点数和轮廓质量敏感。预处理阶段不改变原始文件，只输出到项目的 figures 和 results 目录。")
    add_fig(doc, "fig_q1_coverage.png", "图 2-1 附件 1 候选掩膜覆盖率")
    add_body(doc, "数据观察阶段的关键结论是：图像纹理强度差异大，单一阈值不适合所有样本；Q2 中部分图像不存在稳定的长曲线候选；Q3 轮廓离散化会放大局部斜率；Q4 的图像相似性只能作为连通性弱证据。")

    doc.add_heading("3 模型假设与符号说明", 1)
    add_body(doc, "假设一：经过局部对比度增强和边缘检测后，真实裂隙在候选区域中具有足够的连续性，但不排除纹理和泥浆伪影。假设二：完整正弦状裂隙的周期接近钻孔周长 94.25 mm，周期约束需要通过拟合残差检查。假设三：相邻钻孔同深度段的几何和图像特征相似性可用于构造相对连通评分，但没有现场标注时不把评分解释为真实概率。")
    add_table(doc, ["符号", "含义", "单位"], [["I(x,y)", "灰度图像", "—"], ["R", "正弦裂隙振幅", "mm"], ["P", "正弦裂隙周期", "mm"], ["β", "正弦裂隙相位", "rad"], ["C", "中心线位置", "mm"], ["Z₂", "轮廓均方根斜率", "—"], ["JRC", "裂隙粗糙度系数", "—"], ["p_rel", "相对连通评分", "—"]])

    doc.add_heading("4 问题一 裂隙像素识别", 1)
    doc.add_heading("4.1 问题分析", 2)
    add_body(doc, "裂隙与纹理、泥浆和拼接线在灰度图上可能具有相似的局部边缘，因此问题一不能简单等同于二值阈值分割。本文先建立候选区域生成器，重点保证原图、像素坐标和输出附件之间的一致性；在此基础上，正式方案再引入人工标注和监督分割模型。")
    doc.add_heading("4.2 预处理与候选掩膜", 2)
    add_body(doc, "问题一的目标是生成与原图像素尺寸一致的黑白二值图。本文将局部对比度增强、纹理抑制、暗线增强和边缘检测组合为候选分割链。CLAHE 处理局部亮度不均，双边滤波保留边缘并降低高频纹理，黑帽变换提取相对暗的细长结构，Otsu 给出自适应阈值，Canny 提供梯度边缘，最后通过闭开运算和连通域面积阈值清理噪声。")
    q1p = add_body(doc, "候选掩膜定义为下式，其中 M 的尺寸与输入图像完全一致：")
    add_formula_after(q1p, "M(x,y)=Open(Close(BlackHat(CLAHE(I)) ∪ Canny(I)))")
    add_body(doc, f"10 张图像均成功生成候选掩膜。平均候选覆盖率为 {q1['mean_cov']:.2f}%，覆盖率范围为 {q1['cov_min']:.2f}%—{q1['cov_max']:.2f}%，总连通域数量为 {q1['total_comps']}。覆盖率差异反映纹理和污染差异，不能直接解释为裂隙面积比例。")
    add_fig(doc, "fig_q1_flow.png", "图 4-1 问题一候选分割流程")
    add_fig(doc, "fig_q1_montage.png", "图 4-2 十张附件 1 候选掩膜拼图")
    add_body(doc, "验证方面，当前数据没有独立像素标签，因此无法合法计算 IoU、Dice、Precision、Recall 或 Accuracy。下一步必须建立人工标注协议，按原图分组留出验证集，并以候选分割作为基线比较轻量 U-Net 或 DeepLabV3+。")
    doc.add_heading("4.3 参数敏感性与失败样本", 2)
    add_body(doc, "应对 Canny 高低阈值、黑帽结构元素尺寸、形态学面积下限和局部对比度参数做单因素与联合扰动。若掩膜覆盖率在小幅参数变化下剧烈变化，说明模型对图像质量敏感，需要回到数据增强或学习型分割模型。")
    add_body(doc, "问题小结：Q1 已完成全量候选输出和尺寸检查，但监督验证未完成，当前结果只能支持后续轮廓处理，不能作为最终智能识别结论。")

    doc.add_heading("5 问题二 正弦状裂隙定量分析", 1)
    doc.add_heading("5.1 问题分析与聚类对象", 2)
    add_body(doc, "正弦状裂隙的周期性只对连续裂隙轮廓成立，纹理边缘和拼接线不应被强行拟合。因此先聚类再拟合是必要步骤。DBSCAN 不要求预先给出簇数，并能将稀疏点标记为噪声；其 eps 和 min_samples 必须结合图像尺度、簇跨度和拟合质量检查。")
    doc.add_heading("5.2 坐标标定与参数反演", 2)
    add_body(doc, "问题二首先对附件 2 图像进行高斯平滑和 Canny 边缘提取，再将边缘点按六倍步长降采样，归一化坐标后使用 DBSCAN 聚类。eps=0.055、min_samples=20，点数小于 80 的簇被剔除。为了降低离群点影响，每个簇使用 RANSAC 估计正弦线性化模型。")
    q2p = add_body(doc, "正弦模型写为：")
    add_formula_after(q2p, "y=R sin(2πx/P+β)+C=a sin(2πx/P)+b cos(2πx/P)+C")
    add_body(doc, "固定 P=94.25 mm 后，RANSAC 回归估计 a、b、C，并通过 R=sqrt(a²+b²)、β=atan2(b,a) 恢复参数。每幅图最多保留 R² 最高的三个候选簇。")
    add_body(doc, f"本轮共保留 {q2['n_clusters']} 个候选簇，覆盖 {q2['n_images']} 幅图像，平均 R² 为 {q2['mean_r2']:.4f}，平均 RMSE 为 {q2['mean_rmse']:.2f} mm。最优候选为 {q2['best']['image']} 的簇 {q2['best']['cluster']}，R²={q2['best']['r2']:.4f}。")
    add_fig(doc, "fig_q2_flow.png", "图 5-1 问题二聚类与鲁棒拟合流程")
    add_fig(doc, "fig_q2_r2.png", "图 5-2 各图像候选簇 R² 对比")
    add_fig(doc, "fig_q2_sine.png", "图 5-3 高 R² 候选簇的正弦参数曲线")
    add_body(doc, "结果显示部分图像的候选簇仍然包含纹理和拼接边缘，导致 R² 较低或 RMSE 较大。正式表 1 需要先通过轮廓连续性、周期搜索和人工抽检合并候选簇，再逐裂隙报告 R、P、β、C、RMSE 和 R²。")
    doc.add_heading("5.3 验证、误差来源与问题小结", 2)
    add_body(doc, "误差主要来自边缘点不等于裂隙中心线、DBSCAN 簇可能混入多条结构，以及固定周期假设对局部遮挡不够灵活。正式验证应保存原图、聚类标签、拟合曲线和残差图，并由独立审阅者检查每条曲线是否具有地质意义。")
    add_body(doc, "问题小结：Q2 已建立从聚类到鲁棒拟合的可运行框架，但候选簇质量尚不足以支撑最终裂隙参数表。")

    doc.add_heading("6 问题三 复杂裂隙粗糙度评价", 1)
    doc.add_heading("6.1 轮廓提取与坐标处理", 2)
    add_body(doc, "复杂裂隙的粗糙度评价依赖一条连续、方向正确的轮廓。本文选择面积最大的候选轮廓作为可复现基线，正式方案还需结合位置、长度、连续性和人工复核。重复横坐标先合并，避免斜率计算出现零分母。")
    doc.add_heading("6.2 采样策略与 JRC", 2)
    add_body(doc, "问题三对附件 3 图像进行高斯平滑和 Canny 轮廓提取，筛选长度不小于 30 的轮廓并选面积最大的轮廓作为候选。为了避免相同横坐标引起零分母，先按横坐标合并点并取纵坐标中位数，再将像素坐标映射到钻孔周长和孔深。")
    q3p = add_body(doc, "离散轮廓的粗糙度参数和 JRC 计算为：")
    add_formula_after(q3p, "Z₂=sqrt((1/N)Σ[((yᵢ₊₁−yᵢ)/(xᵢ₊₁−xᵢ))²]),   JRC=51.85Z₂^0.6−10.37")
    add_body(doc, f"11 张图像均提取到候选轮廓，共生成 {q3['n_schemes']} 组采样结果。所有方案的 JRC 范围为 {q3['range_all'][0]:.2f}—{q3['range_all'][1]:.2f}，均匀采样均值为 {q3['mean_uniform']:.2f}，曲率自适应采样均值为 {q3['mean_adaptive']:.2f}。")
    add_fig(doc, "fig_q3_flow.png", "图 6-1 问题三轮廓与 JRC 计算流程")
    add_fig(doc, "fig_q3_sens.png", "图 6-2 不同采样方案的 JRC 敏感性")
    add_body(doc, "采样点数增大并不一定提高可靠性，因为边缘噪声会被斜率平方放大。正式方案应报告轮廓平滑参数、曲率加权规则、面积分层结果和异常轮廓，并把 JRC 与 Barton 标准线或工程标注进行校准。")
    add_body(doc, "采样点数增大不一定提高可靠性，因为 JRC 使用斜率平方，噪声尖峰会被放大。正式方案需要联合分析平滑尺度、面积和采样密度。")
    doc.add_heading("6.3 验证与问题小结", 2)
    add_body(doc, "若无法获得 Barton 标准轮廓或现场粗糙度标定，应将 JRC 表述为图像派生的相对粗糙度，而不是工程等级。问题小结：Q3 已完成轮廓提取和采样比较，但工程标定和面积效应仍需补充。")

    doc.add_heading("7 问题四 多钻孔连通性与三维重构", 1)
    doc.add_heading("7.1 空间布置与匹配对象", 2)
    add_body(doc, "六个钻孔按 2×3 阵列布置，孔间距为 1000 mm。附件 4 每张图对应 1000 mm 孔深，4 号孔深度较短。匹配时必须区分孔号、分段起点和图像方向，不能仅按文件名排序。")
    doc.add_heading("7.2 相对评分与不确定性", 2)
    add_body(doc, "问题四将六个钻孔的空间坐标统一到同一坐标系，对每个分段计算边缘密度、平均亮度和竖直梯度，并对七组相邻钻孔对的同深度段进行匹配。设两段图像的特征差异为 d，构造相对连通评分：")
    q4p = add_body(doc, "")
    add_formula_after(q4p, "p_rel=exp(−50d),   d=|Δρ_edge|+0.002|ΔI|/255")
    add_body(doc, f"七组相邻孔对均生成了评分，范围为 {q4['range'][0]:.4f}—{q4['range'][1]:.4f}，均值为 {q4['mean_score']:.4f}。按最接近 0.5 的不确定性排序，前三个候选为 {q4['unc'][0]['pair']}、{q4['unc'][1]['pair']} 和 {q4['unc'][2]['pair']}。")
    add_fig(doc, "fig_q4_flow.png", "图 7-1 问题四连通评分与补孔流程")
    add_fig(doc, "fig_q4_scores.png", "图 7-2 相邻钻孔对相对连通评分")
    add_fig(doc, "fig_q4_3d.png", "图 7-3 六钻孔空间布置与相对连通网络")
    add_body(doc, "当前评分没有整合裂隙方向、相位、JRC、三维姿态和现场连通标签，因而只能作为相对证据。正式模型应将裂隙片段作为节点、几何匹配作为候选边，并用标定数据或重复情景建立概率校准；补孔优化还需加入工程成本、施工可达性和覆盖收益约束。")
    add_body(doc, "相对评分只能识别需要进一步观测的孔对。正式模型应引入裂隙方向、相位、JRC、三维姿态和现场连通标签，并用可靠性曲线或 Brier 分数校准概率。问题小结：Q4 已完成数据链路、相邻孔评分和三维可视化，但尚未形成可用于工程决策的概率模型。")

    doc.add_heading("8 模型验证、敏感性与工程解释", 1)
    add_body(doc, "四问的验证必须与任务类型匹配。Q1 需要像素标签和独立验证；Q2 需要逐裂隙拟合残差、R²、RMSE 和参数稳定性；Q3 需要采样密度、平滑参数和面积效应敏感性；Q4 需要概率校准、几何回代误差、阈值扰动和补孔收益/成本比较。当前已实际完成的验证主要是输入覆盖、尺寸一致、重复文件审计、RANSAC 拟合统计和采样敏感性，尚未满足正式交稿标准。")
    add_table(doc, ["问题", "已完成验证", "仍需补齐"], [["Q1", "尺寸、覆盖率、候选连通域", "独立标签、IoU/Dice/F1、失败样本"], ["Q2", "R²、RMSE、内点比例", "逐裂隙聚类、周期搜索、人工复核"], ["Q3", "采样点数敏感性", "平滑/面积敏感性、工程标定"], ["Q4", "评分范围、不确定性排序", "概率校准、几何回代、成本约束"]])
    add_body(doc, "工程解释必须遵守结论边界：候选掩膜只表示算法认为可能存在裂隙的区域；低 R² 表示正弦假设或轮廓分离不充分；JRC 变化表示采样和轮廓误差敏感；相对连通评分只能用于安排进一步检测优先级。")

    doc.add_heading("9 模型评价、创新与推广", 1)
    add_body(doc, "本文的可复现性体现在输入哈希、代码入口、参数和 CSV 结果之间的对应关系。方法上的初步创新包括把四问组织为一条裂隙信息链、用候选簇的 RANSAC 内点统计过滤异常边缘、在 JRC 计算前合并重复横坐标、以及把跨孔评分转化为不确定性驱动的补孔候选。由于缺少标注和现场校准，这些内容目前属于可验证的方案设计，不应写成已经证明有效的创新结论。")
    add_body(doc, "推广方向包括：使用带标注的 U-Net/DeepLabV3+ 完成像素分割；使用 DBSCAN、GMM 和几何连续性完成逐裂隙聚类；使用 RANSAC、Huber 或非线性最小二乘提高正弦拟合鲁棒性；使用曲率自适应采样和标准轮廓校准 JRC；使用贝叶斯图模型或校准分类器估计连通概率，并用整数规划选择补孔方案。")

    doc.add_heading("参考文献", 1)
    for ref in [
        "[1] 康红普，姜鹏飞，王子越，等. 煤巷钻锚一体化快速掘进技术与装备及应用[J]. 煤炭学报，2024，49(1):131-151。",
        "[2] 袁亮，张平松. 煤矿透明地质模型动态重构的关键技术与路径思考[J]. 煤炭学报，2023，48(1):1-14。",
        "[3] Canny J. A computational approach to edge detection[J]. IEEE TPAMI, 1986, 8(6):679-698。",
        "[4] Ronneberger O, Fischer P, Brox T. U-Net: Convolutional networks for biomedical image segmentation[C]. MICCAI, 2015。",
        "[5] Ester M, Kriegel H-P, Sander J, et al. A density-based algorithm for discovering clusters in large spatial databases[C]. KDD, 1996。",
        "[6] Fischler M A, Bolles R C. Random sample consensus[J]. Communications of the ACM, 1981, 24(6):381-395。",
        "[7] Barton N, Choubey V. The shear strength of rock joints in theory and practice[J]. Rock Mechanics, 1977, 10:1-54。",
        "[8] 2025 年中国研究生数学建模竞赛 C 题题面及附件说明，竞赛资料目录提供的原始文件。",
    ]:
        add_body(doc, ref)

    doc.add_heading("附录 代码与复现说明", 1)
    add_body(doc, "项目使用 uv 管理 Python 环境，依赖锁定在项目根目录 uv.lock。四问代码入口分别为 projects/2025-C/code/q1_segmentation.py、q2_sine_fit.py、q3_jrc.py 和 q4_connectivity.py；图表入口为 make_project_figures.py 与 make_paper_figures.py；结果汇总为 results/paper_data.json。")
    add_body(doc, "本稿使用真实附件和实际运行结果，但尚未通过国奖级参赛门禁。正式候选稿必须补齐当届官方规则快照、Q1 独立标注验证、Q2 逐裂隙表征、Q3 工程 JRC 标定、Q4 概率校准、证据矩阵、独立技术审查和独立论文审查。")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
